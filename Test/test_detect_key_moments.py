from __future__ import annotations

import datetime as dt
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "Utilities"))
import detect_key_moments as detect


class AnnotationWorkbookTests(unittest.TestCase):
    def make_book(self, folder: Path, rows, headers=("Time (seconds)", "Label"), *, sheet="KeyMoments"):
        book = Workbook()
        tab = book.active
        tab.title = sheet
        tab.append(list(headers))
        for row in rows:
            tab.append(row)
        path = folder / "examples.xlsx"
        book.save(path)
        book.close()
        return path

    def test_numeric_text_excel_duration_and_blank_rows(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp:
            path = self.make_book(Path(temp), [[2.25, "goal"], ["00:00:04.500", "play"],
                                                [dt.timedelta(seconds=6), "save"], [None, None]])
            entries = detect.read_annotations(path, 10)
        self.assertEqual([entry.time_seconds for entry in entries], [2.25, 4.5, 6.0])
        self.assertEqual(entries[1].row_number, 3)

    def test_rejects_missing_sheet_and_headers(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp:
            folder = Path(temp)
            missing_sheet = self.make_book(folder, [[1, "x"]], sheet="Other")
            with self.assertRaisesRegex(detect.DetectionError, "KeyMoments"):
                detect.read_annotations(missing_sheet, 10)
            wrong_header = self.make_book(folder, [], headers=("Time", "Name"))
            with self.assertRaisesRegex(detect.DetectionError, "headers"):
                detect.read_annotations(wrong_header, 10)

    def test_rejects_empty_missing_labels_bad_times_out_of_range_duplicates_and_formulas(self):
        cases = [([], "no valid examples"), ([[1, None]], "Label"),
                 ([[None, "goal"]], "nonnegative time"), ([[-1, "goal"]], "nonnegative"),
                 ([["00:61:00", "goal"]], "HH:MM:SS"), ([["1:2:03", "goal"]], "HH:MM:SS"),
                 ([[11, "goal"]], "exceeds clip"),
                 ([[1, "goal"], [1.0, "other"]], "duplicate")]
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            for rows, error in cases:
                with self.subTest(rows=rows):
                    path = self.make_book(folder, rows)
                    with self.assertRaisesRegex(detect.DetectionError, error):
                        detect.read_annotations(path, 10)
            formula = self.make_book(folder, [["=1+1", "goal"]])
            with self.assertRaisesRegex(detect.DetectionError, "formulas"):
                detect.read_annotations(formula, 10)


class CandidateAndOutputTests(unittest.TestCase):
    def test_candidates_suppress_known_times_and_merge_close_peaks(self):
        profile = detect.SPORT_PROFILES[detect.SportType.SOCCER]
        annotations = (detect.Annotation(5, "known", 2),)
        scores = [(float(t), s) for t, s in [(1, .1), (2, .5), (3, .1), (5, .8), (6, .2),
                                              (9, .3), (10, .7), (11, .3), (14, .4)]]
        found = detect._select_candidates(scores, annotations, profile, .18, 15)
        self.assertEqual([round(row.time_seconds) for row in found], [2, 10, 14])
        self.assertEqual([round(row.score, 1) for row in found], [.5, .7, .4])

    def test_output_contract_empty_and_nonempty_and_existing_behavior(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp:
            path = Path(temp) / "result.xlsx"
            detect._write_workbook(path, (), False)
            book = load_workbook(path, read_only=True, data_only=False)
            sheet = book["Detected Moments"]
            self.assertEqual([cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))],
                             ["Time (seconds)", "Confidence (0-1)"])
            self.assertEqual(sheet.max_row, 1)
            book.close()
            with self.assertRaisesRegex(detect.DetectionError, "already exists"):
                detect._write_workbook(path, (), False)
            detect._write_workbook(path, (detect.Detection(2.12345, .75),), True)
            book = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(book["Detected Moments"]["A2"].value, 2.123)
            self.assertEqual(book["Detected Moments"]["B2"].value, .75)
            book.close()

    def test_profiles_are_distinct_and_sport_input_is_enum_constrained(self):
        self.assertEqual(set(detect.SPORT_PROFILES), set(detect.SportType))
        self.assertNotEqual(detect.SPORT_PROFILES[detect.SportType.SOCCER],
                            detect.SPORT_PROFILES[detect.SportType.HOCKEY])
        with self.assertRaisesRegex(detect.DetectionError, "Supported values"):
            detect.detect_key_moments("missing.mp4", "tennis", "missing.xlsx")

    def test_success_writes_unlabeled_candidates_and_leaves_inputs_unchanged(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp:
            folder = Path(temp)
            video = folder / "trimmed.mp4"
            video.write_bytes(b"video bytes")
            examples = AnnotationWorkbookTests().make_book(folder, [[1.0, "known goal"]])
            output = folder / "found.xlsx"
            video_hash = hashlib.sha256(video.read_bytes()).hexdigest()
            examples_hash = hashlib.sha256(examples.read_bytes()).hexdigest()
            info = detect.VideoInfo(5.0, 15.0, 1280, 720)
            with patch.object(detect, "_probe_video", return_value=info):
                with patch.object(detect, "_scan", return_value=([(1.0, .4), (3.0, .6)], .2)):
                    result = detect.detect_key_moments(video, detect.SportType.SOCCER, examples, output)
            self.assertEqual(result.detections, (detect.Detection(3.0, .6),))
            self.assertEqual(hashlib.sha256(video.read_bytes()).hexdigest(), video_hash)
            self.assertEqual(hashlib.sha256(examples.read_bytes()).hexdigest(), examples_hash)
            book = load_workbook(output, read_only=True, data_only=True)
            sheet = book["Detected Moments"]
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual([cell.value for cell in sheet[2]], [3, .6])
            self.assertNotIn("Label", [cell.value for cell in sheet[1]])
            book.close()

    def test_unavailable_opencv_has_install_instruction(self):
        with patch.dict(sys.modules, {"cv2": None}):
            with self.assertRaisesRegex(detect.DetectionError, "OpenCV.*pip install"):
                detect._opencv()

    def test_malformed_or_unreadable_video_is_reported(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp:
            video = Path(temp) / "empty.mp4"
            video.touch()
            with patch.object(detect.shutil, "which", return_value=None):
                with self.assertRaisesRegex(detect.DetectionError, "ffprobe"):
                    detect._probe_video(video)


if __name__ == "__main__":
    unittest.main()
