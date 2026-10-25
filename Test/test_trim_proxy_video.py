from __future__ import annotations

import datetime as dt
import json
import os
import shutil
import subprocess
import sys
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "Utilities"))

import trim_proxy_video as trim


@contextmanager
def test_file_prefix():
    prefix = Path(__file__).parent / f".trim-test-{uuid.uuid4().hex}"
    try:
        yield prefix
    finally:
        for path in Path(__file__).parent.iterdir():
            if path.is_file() and (
                path.name.startswith(prefix.name) or path.name.startswith(f".{prefix.name}")
            ):
                path.unlink(missing_ok=True)


class WorkbookParsingTests(unittest.TestCase):
    def make_workbook(self, prefix: Path, rows, headers=None) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Segments"
        sheet.append(headers or ["Start Time (seconds)", "Stop Time (seconds)", "Label"])
        for row in rows:
            sheet.append(row)
        path = prefix.with_name(prefix.name + "_ranges.xlsx")
        workbook.save(path)
        workbook.close()
        return path

    def test_numeric_seconds_rows_and_blank_rows(self):
        with test_file_prefix() as prefix:
            path = self.make_workbook(prefix, [[0, 1.1, "opening"], [None, None, None], [2, 3, "play"]])
            segments = trim.read_segments(path, duration_seconds=10, fps=10)
        self.assertEqual(len(segments), 2)
        self.assertEqual(segments[0].start_frame, 0)
        self.assertEqual(segments[0].stop_frame, 11)
        self.assertEqual(segments[1].label, "play")

    def test_elapsed_time_text_and_excel_duration(self):
        with test_file_prefix() as prefix:
            path = self.make_workbook(
                prefix,
                [["00:00:01.250", "00:00:02.500", "text"],
                 [dt.timedelta(seconds=3), dt.timedelta(seconds=4), "excel duration"]],
            )
            segments = trim.read_segments(path, duration_seconds=10, fps=20)
        self.assertAlmostEqual(segments[0].start_seconds, 1.25)
        self.assertAlmostEqual(segments[1].stop_seconds, 4.0)

    def test_rejects_missing_sheet_and_headers(self):
        with test_file_prefix() as prefix:
            no_sheet = Workbook()
            no_sheet.active.title = "Other"
            no_sheet_path = prefix.with_name(prefix.name + "_no_sheet.xlsx")
            no_sheet.save(no_sheet_path)
            no_sheet.close()
            with self.assertRaisesRegex(trim.TrimError, "Segments"):
                trim.read_segments(no_sheet_path, 10, 10)

            wrong_headers = self.make_workbook(prefix, [], headers=["Start", "Stop"])
            with self.assertRaisesRegex(trim.TrimError, "headers"):
                trim.read_segments(wrong_headers, 10, 10)

    def test_rejects_formula_and_partially_filled_rows(self):
        with test_file_prefix() as prefix:
            formula = self.make_workbook(prefix, [["=1", 2, "formula"]])
            with self.assertRaisesRegex(trim.TrimError, "formulas"):
                trim.read_segments(formula, 10, 10)
            partial = self.make_workbook(prefix, [[1, None, "missing stop"]])
            with self.assertRaisesRegex(trim.TrimError, "Row 2"):
                trim.read_segments(partial, 10, 10)

    def test_rejects_empty_and_invalid_ranges(self):
        invalid_cases = [
            ([], "No clip ranges"),
            ([[-1, 2, "negative"]], "nonnegative"),
            ([[3, 3, "zero length"]], "later than"),
            ([[0, 10.01, "too long"]], "exceeds proxy duration"),
            ([[1, 3, "first"], [2, 4, "overlap"]], "chronological and non-overlapping"),
            ([[3, 4, "later"], [1, 2, "out of order"]], "chronological and non-overlapping"),
            ([[1.001, 1.02, "less than one frame"]], "shorter than one frame"),
            ([["not-a-time", 2, "invalid"]], "decimal seconds"),
        ]
        with test_file_prefix() as prefix:
            for rows, message in invalid_cases:
                with self.subTest(rows=rows):
                    path = self.make_workbook(prefix, rows)
                    with self.assertRaisesRegex(trim.TrimError, message):
                        trim.read_segments(path, duration_seconds=10, fps=10)

    def test_rounding_is_nearest_frame_with_half_up_ties(self):
        with test_file_prefix() as prefix:
            path = self.make_workbook(prefix, [[0.05, 0.25, "ties"]])
            segment = trim.read_segments(path, duration_seconds=10, fps=10)[0]
        self.assertEqual(segment.start_frame, 1)
        self.assertEqual(segment.stop_frame, 3)


class TrimWorkflowTests(unittest.TestCase):
    def make_inputs(self, prefix: Path, rows):
        source = prefix.with_name(prefix.name + "_proxy.mp4")
        source.write_bytes(b"source proxy")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Segments"
        sheet.append(["Start Time (seconds)", "Stop Time (seconds)", "Label"])
        for row in rows:
            sheet.append(row)
        ranges = prefix.with_name(prefix.name + "_ranges.xlsx")
        workbook.save(ranges)
        workbook.close()
        return source, ranges

    def test_filter_graph_joins_video_and_audio(self):
        segments = (
            trim.Segment(0, 1, "a", 2, 0, 15, 15),
            trim.Segment(2, 3, "b", 3, 30, 45, 15),
        )
        graph = trim._filter_graph(segments, with_audio=True)
        self.assertIn("split=2", graph)
        self.assertIn("trim=start_frame=0:end_frame=15", graph)
        self.assertIn("concat=n=2:v=1:a=0[vout]", graph)
        self.assertIn("concat=n=2:v=0:a=1[aout]", graph)

    def test_notebook_finds_repository_root_from_root_and_test(self):
        root = Path(__file__).resolve().parents[1]
        notebook_path = root / "Test" / "trim_proxy_video_example.ipynb"
        notebook = json.loads(notebook_path.read_text(encoding="utf-8"))
        root_cell = next(
            cell for cell in notebook["cells"]
            if cell["cell_type"] == "code" and any("starting_directory" in line for line in cell["source"])
        )
        original_directory = Path.cwd()
        try:
            for working_directory in (root, root / "Test"):
                os.chdir(working_directory)
                scope: dict[str, object] = {}
                exec("".join(root_cell["source"]), scope)
                self.assertEqual(scope["repo_root"], root)
        finally:
            os.chdir(original_directory)

    def test_missing_ffmpeg_is_reported_without_output(self):
        with test_file_prefix() as prefix:
            source, ranges = self.make_inputs(prefix, [[0, 1, "one"]])
            output = prefix.with_name(prefix.name + "_out.mp4")
            with patch.object(trim.shutil, "which", return_value=None):
                with self.assertRaisesRegex(trim.TrimError, "ffmpeg, ffprobe"):
                    trim.trim_proxy(source, ranges, output)
            self.assertFalse(output.exists())

    def test_success_publishes_result_and_leaves_inputs_unchanged(self):
        fake_info = trim.VideoInfo(1280, 720, 1280, 720, 15.0, 5.0, True, "h264")
        with test_file_prefix() as prefix:
            source, ranges = self.make_inputs(prefix, [[0.04, 1.04, "keep"]])
            output = prefix.with_name(prefix.name + "_out.mp4")
            source_before = source.read_bytes()
            ranges_before = ranges.read_bytes()

            def fake_run(command, action):
                Path(command[-1]).write_bytes(b"trimmed output")

            with patch.object(trim.shutil, "which", side_effect=lambda name: name):
                with patch.object(trim, "_probe_video", side_effect=[fake_info, fake_info]):
                    with patch.object(trim, "_run", side_effect=fake_run):
                        result = trim.trim_proxy(source, ranges, output)

            self.assertEqual(output.read_bytes(), b"trimmed output")
            self.assertEqual(source.read_bytes(), source_before)
            self.assertEqual(ranges.read_bytes(), ranges_before)
            self.assertAlmostEqual(result.kept_duration_seconds, 1.0)
            self.assertAlmostEqual(result.segments[0].effective_start_seconds, 1 / 15)

    def test_failed_encode_leaves_no_final_or_temporary_output(self):
        fake_info = trim.VideoInfo(1280, 720, 1280, 720, 15.0, 5.0, False, "h264")
        with test_file_prefix() as prefix:
            source, ranges = self.make_inputs(prefix, [[0, 1, "keep"]])
            output = prefix.with_name(prefix.name + "_out.mp4")

            def fail_after_partial(command, action):
                Path(command[-1]).write_bytes(b"partial")
                raise trim.TrimError("synthetic encode failure")

            with patch.object(trim.shutil, "which", side_effect=lambda name: name):
                with patch.object(trim, "_probe_video", return_value=fake_info):
                    with patch.object(trim, "_run", side_effect=fail_after_partial):
                        with self.assertRaisesRegex(trim.TrimError, "synthetic encode failure"):
                            trim.trim_proxy(source, ranges, output)
            self.assertFalse(output.exists())
            self.assertEqual(list(prefix.parent.glob(f".{prefix.name}_out-*.tmp.mp4")), [])

    def test_existing_output_and_same_input_output_are_rejected(self):
        with test_file_prefix() as prefix:
            source, ranges = self.make_inputs(prefix, [[0, 1, "keep"]])
            output = prefix.with_name(prefix.name + "_out.mp4")
            output.write_bytes(b"keep existing")
            with self.assertRaisesRegex(trim.TrimError, "already exists"):
                trim.trim_proxy(source, ranges, output)
            with self.assertRaisesRegex(trim.TrimError, "three distinct files"):
                trim.trim_proxy(source, ranges, source)
            self.assertEqual(output.read_bytes(), b"keep existing")

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg tools unavailable")
    def test_real_proxy_trim_and_audio_modes(self):
        root = Path(__file__).resolve().parents[1]
        source = root / "output" / "proxies" / "26_10_24_GoPro_Goal_proxy.mp4"
        if not source.is_file():
            self.skipTest("Step 1 proxy is not present in this checkout")
        with test_file_prefix() as prefix:
            ranges = prefix.with_name(prefix.name + "_ranges.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Segments"
            sheet.append(["Start Time (seconds)", "Stop Time (seconds)", "Label"])
            sheet.append([1, 4, "first clip"])
            sheet.append([8, 11, "second clip"])
            workbook.save(ranges)
            workbook.close()
            ranges_before = ranges.read_bytes()
            source_size = source.stat().st_size

            single_ranges = prefix.with_name(prefix.name + "_single_ranges.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Segments"
            sheet.append(["Start Time (seconds)", "Stop Time (seconds)", "Label"])
            sheet.append([3, 6, "single clip"])
            workbook.save(single_ranges)
            workbook.close()
            single_output = prefix.with_name(prefix.name + "_single.mp4")
            single = trim.trim_proxy(source, single_ranges, single_output)
            self.assertEqual(len(single.segments), 1)
            self.assertAlmostEqual(single.output.duration_seconds, 3, delta=0.2)

            with_audio = prefix.with_name(prefix.name + "_with_audio.mp4")
            kept = trim.trim_proxy(source, ranges, with_audio)
            self.assertEqual(len(kept.segments), 2)
            self.assertTrue(kept.output.has_audio)
            self.assertEqual(kept.output.display_width, 1280)
            self.assertEqual(kept.output.display_height, 720)
            self.assertAlmostEqual(kept.output.frame_rate, 15, delta=0.02)
            self.assertAlmostEqual(kept.output.duration_seconds, 6, delta=0.2)

            silent = prefix.with_name(prefix.name + "_silent_source.mp4")
            ffmpeg = shutil.which("ffmpeg")
            self.assertIsNotNone(ffmpeg)
            subprocess.run(
                [ffmpeg, "-hide_banner", "-v", "error", "-nostdin", "-y", "-i", str(source),
                 "-map", "0:v:0", "-c:v", "copy", "-an", str(silent)],
                check=True, capture_output=True,
            )
            without_audio = prefix.with_name(prefix.name + "_without_audio.mp4")
            removed = trim.trim_proxy(source, ranges, without_audio, include_audio=False)
            self.assertFalse(removed.output.has_audio)
            silent_trimmed = prefix.with_name(prefix.name + "_silent_trim.mp4")
            no_audio_source = trim.trim_proxy(silent, ranges, silent_trimmed, include_audio=True)
            self.assertFalse(no_audio_source.source.has_audio)
            self.assertFalse(no_audio_source.output.has_audio)

            self.assertEqual(source.stat().st_size, source_size)
            self.assertEqual(ranges.read_bytes(), ranges_before)


if __name__ == "__main__":
    unittest.main()
