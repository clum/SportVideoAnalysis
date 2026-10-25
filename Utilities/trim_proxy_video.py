"""Trim a proxy video to spreadsheet-selected time ranges.

Requires the Python packages in ``Utilities/requirements.txt`` and FFmpeg plus
ffprobe on PATH. Spreadsheet times are elapsed seconds from the start of the
proxy. Each interval is start-inclusive and stop-exclusive; its boundaries are
rounded to the nearest source frame (half-frame ties round forward).

Python API::

    result = trim_proxy(
        "output/proxies/source_proxy.mp4",
        "Test/segments_template.xlsx",
        "output/clips/selected.mp4",
    )

Command line::

    python Utilities/trim_proxy_video.py proxy.mp4 segments.xlsx selected.mp4
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any


class TrimError(RuntimeError):
    """Raised when the workbook, source video, or trimming operation is invalid."""


@dataclass(frozen=True)
class Segment:
    """One requested range and its effective frame-aligned boundaries."""

    start_seconds: float
    stop_seconds: float
    label: str
    row_number: int
    start_frame: int
    stop_frame: int
    fps: float

    @property
    def effective_start_seconds(self) -> float:
        return self.start_frame / self.fps

    @property
    def effective_stop_seconds(self) -> float:
        return self.stop_frame / self.fps

    @property
    def frame_count(self) -> int:
        return self.stop_frame - self.start_frame


@dataclass(frozen=True)
class VideoInfo:
    width: int
    height: int
    display_width: int
    display_height: int
    frame_rate: float
    duration_seconds: float
    has_audio: bool
    video_codec: str | None


@dataclass(frozen=True)
class TrimResult:
    source: VideoInfo
    output: VideoInfo
    segments: tuple[Segment, ...]
    output_path: Path
    kept_duration_seconds: float
    audio_included: bool


_TIME_TEXT = re.compile(r"^(\d+):(\d{2}):(\d{2})(?:\.(\d{1,6}))?$")


def _as_seconds(value: Any, row_number: int, column_name: str) -> float:
    if isinstance(value, bool) or value is None:
        raise TrimError(f"Row {row_number}: {column_name} must be a nonnegative time value.")

    if isinstance(value, dt.timedelta):
        seconds = value.total_seconds()
    elif isinstance(value, dt.time):
        seconds = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    elif isinstance(value, (int, float)):
        seconds = float(value)
    elif isinstance(value, str):
        text = value.strip()
        match = _TIME_TEXT.fullmatch(text)
        if match:
            hours, minutes, whole_seconds, fraction = match.groups()
            sec_value = int(whole_seconds)
            min_value = int(minutes)
            if min_value >= 60 or sec_value >= 60:
                raise TrimError(
                    f"Row {row_number}: {column_name} time text must use HH:MM:SS[.ffffff] format."
                )
            fractional_seconds = float(f"0.{fraction}") if fraction else 0.0
            seconds = int(hours) * 3600 + min_value * 60 + sec_value + fractional_seconds
        else:
            try:
                seconds = float(text)
            except ValueError as exc:
                raise TrimError(
                    f"Row {row_number}: {column_name} must be decimal seconds or HH:MM:SS[.mmm] text."
                ) from exc
    else:
        raise TrimError(
            f"Row {row_number}: {column_name} must be decimal seconds or an Excel elapsed-time value."
        )

    if not math.isfinite(seconds) or seconds < 0:
        raise TrimError(f"Row {row_number}: {column_name} must be a finite, nonnegative time.")
    return float(seconds)


def _nearest_frame(seconds: float, fps: float) -> int:
    # Times are nonnegative, so floor(x + 0.5) gives half-up rounding.
    return math.floor(seconds * fps + 0.5)


def read_segments(workbook_path: str | Path, duration_seconds: float, fps: float) -> tuple[Segment, ...]:
    """Read, validate, and frame-align ranges from the Segments worksheet."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise TrimError(
            "Reading .xlsx files requires openpyxl. Install project dependencies with "
            "'python -m pip install -r Utilities/requirements.txt'."
        ) from exc

    path = Path(workbook_path).expanduser()
    if not path.is_file():
        raise TrimError(f"Range workbook does not exist or is not a file: '{path}'.")
    if path.suffix.lower() != ".xlsx":
        raise TrimError("Range workbook must be an .xlsx file.")
    if not math.isfinite(duration_seconds) or duration_seconds <= 0:
        raise TrimError("The proxy must have a known, positive duration before ranges can be validated.")
    if not math.isfinite(fps) or fps <= 0:
        raise TrimError("The proxy must have a known, positive frame rate before ranges can be validated.")

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise TrimError(f"Could not read .xlsx workbook '{path}': {exc}") from exc
    try:
        if "Segments" not in workbook.sheetnames:
            raise TrimError("Workbook must contain a worksheet named 'Segments'.")
        sheet = workbook["Segments"]
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers: dict[str, int] = {}
        for index, cell in enumerate(header_cells):
            value = cell.value
            if isinstance(value, str):
                headers[value.strip().casefold()] = index
        start_key = "start time (seconds)"
        stop_key = "stop time (seconds)"
        if start_key not in headers or stop_key not in headers:
            raise TrimError(
                "The 'Segments' worksheet must have 'Start Time (seconds)' and "
                "'Stop Time (seconds)' headers in row 1."
            )
        label_index = headers.get("label")
        segments: list[Segment] = []
        previous_start = -1.0
        previous_stop = -1.0
        previous_effective_stop_frame = -1
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if all(cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()) for cell in row):
                continue
            for index in (headers[start_key], headers[stop_key], label_index):
                if index is not None and index < len(row) and row[index].data_type == "f":
                    raise TrimError(f"Row {row_number}: formulas are not allowed in the Segments input cells.")

            start_cell = row[headers[start_key]]
            stop_cell = row[headers[stop_key]]
            label = ""
            if label_index is not None and label_index < len(row) and row[label_index].value is not None:
                label = str(row[label_index].value).strip()
            if start_cell.value is None or stop_cell.value is None:
                raise TrimError(
                    f"Row {row_number}: enter both Start Time and Stop Time, or leave the entire row blank."
                )

            start = _as_seconds(start_cell.value, row_number, "Start Time")
            stop = _as_seconds(stop_cell.value, row_number, "Stop Time")
            if start >= stop:
                raise TrimError(f"Row {row_number}: Stop Time must be later than Start Time.")
            if stop > duration_seconds + 1e-9:
                raise TrimError(
                    f"Row {row_number}: Stop Time ({stop:g} s) exceeds proxy duration ({duration_seconds:g} s)."
                )
            if start + 1e-9 < previous_start or start + 1e-9 < previous_stop:
                raise TrimError(
                    f"Row {row_number}: ranges must be chronological and non-overlapping "
                    "in worksheet row order."
                )

            start_frame = _nearest_frame(start, fps)
            stop_frame = _nearest_frame(stop, fps)
            if stop_frame <= start_frame:
                raise TrimError(
                    f"Row {row_number}: range is shorter than one frame after nearest-frame rounding."
                )
            if start_frame < previous_effective_stop_frame:
                raise TrimError(
                    f"Row {row_number}: frame rounding makes this range overlap the prior range."
                )
            segments.append(
                Segment(start, stop, label, row_number, start_frame, stop_frame, fps)
            )
            previous_start = start
            previous_stop = stop
            previous_effective_stop_frame = stop_frame

        if not segments:
            raise TrimError("No clip ranges were found. Enter at least one range on the Segments worksheet.")
        return tuple(segments)
    finally:
        workbook.close()


def _parse_rate(value: Any) -> float | None:
    if value in (None, "", "N/A", "0/0"):
        return None
    try:
        if isinstance(value, str) and "/" in value:
            numerator, denominator = value.split("/", maxsplit=1)
            result = float(numerator) / float(denominator)
        else:
            result = float(value)
    except (TypeError, ValueError, ZeroDivisionError):
        return None
    return result if math.isfinite(result) else None


def _probe_video(path: Path, ffprobe: str) -> VideoInfo:
    command = [
        ffprobe,
        "-v", "error",
        "-show_entries",
        "stream=codec_type,codec_name,width,height,avg_frame_rate,r_frame_rate,duration:stream_tags=rotate:stream_side_data=rotation:format=duration",
        "-of", "json",
        str(path),
    ]
    try:
        completed = subprocess.run(
            command, check=True, capture_output=True, text=True, encoding="utf-8", errors="replace"
        )
    except FileNotFoundError as exc:
        raise TrimError("ffprobe was not found. Install FFmpeg and add ffprobe to PATH.") from exc
    except subprocess.CalledProcessError as exc:
        detail = (exc.stderr or "").strip()
        raise TrimError(f"Could not probe proxy video '{path}'. {detail}".strip()) from exc
    except OSError as exc:
        raise TrimError(f"Could not run ffprobe: {exc}") from exc

    try:
        metadata = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise TrimError("ffprobe returned invalid JSON for the proxy video.") from exc
    streams = metadata.get("streams") or []
    video = next((s for s in streams if s.get("codec_type") == "video"), None)
    if video is None:
        raise TrimError(f"No video stream was found in '{path}'.")
    try:
        width, height = int(video["width"]), int(video["height"])
    except (KeyError, TypeError, ValueError) as exc:
        raise TrimError("Proxy video dimensions are missing or invalid.") from exc
    fps = _parse_rate(video.get("avg_frame_rate")) or _parse_rate(video.get("r_frame_rate"))
    duration = _parse_rate(video.get("duration"))
    if duration is None:
        duration = _parse_rate((metadata.get("format") or {}).get("duration"))
    if not fps or fps <= 0:
        raise TrimError("The proxy's frame rate is missing or invalid; cannot align range boundaries.")
    if duration is None or duration <= 0:
        raise TrimError("The proxy's duration is missing or invalid; cannot validate time ranges.")

    rotation = _parse_rate((video.get("tags") or {}).get("rotate"))
    for side_data in video.get("side_data_list") or []:
        side_rotation = _parse_rate(side_data.get("rotation"))
        if side_rotation is not None:
            rotation = side_rotation
            break
    display_width, display_height = (height, width) if int(round(rotation or 0)) % 180 else (width, height)
    return VideoInfo(
        width=width,
        height=height,
        display_width=display_width,
        display_height=display_height,
        frame_rate=fps,
        duration_seconds=duration,
        has_audio=any(s.get("codec_type") == "audio" for s in streams),
        video_codec=video.get("codec_name"),
    )


def _filter_graph(segments: tuple[Segment, ...], with_audio: bool) -> str:
    count = len(segments)
    graph: list[str] = []
    video_inputs: list[str] = []
    if count > 1:
        outputs = "".join(f"[vsrc{i}]" for i in range(count))
        graph.append(f"[0:v:0]split={count}{outputs}")
    for i, segment in enumerate(segments):
        label = f"vsrc{i}" if count > 1 else "0:v:0"
        graph.append(
            f"[{label}]trim=start_frame={segment.start_frame}:end_frame={segment.stop_frame},"
            f"setpts=PTS-STARTPTS[v{i}]"
        )
        video_inputs.append(f"[v{i}]")
    graph.append(f"{''.join(video_inputs)}concat=n={count}:v=1:a=0[vout]")

    if with_audio:
        if count > 1:
            outputs = "".join(f"[asrc{i}]" for i in range(count))
            graph.append(f"[0:a:0]asplit={count}{outputs}")
        audio_inputs: list[str] = []
        for i, segment in enumerate(segments):
            label = f"asrc{i}" if count > 1 else "0:a:0"
            start = segment.effective_start_seconds
            stop = segment.effective_stop_seconds
            graph.append(
                f"[{label}]atrim=start={start:.9f}:end={stop:.9f},"
                f"asetpts=PTS-STARTPTS[a{i}]"
            )
            audio_inputs.append(f"[a{i}]")
        graph.append(f"{''.join(audio_inputs)}concat=n={count}:v=0:a=1[aout]")
    return ";".join(graph)


def _run(command: list[str], action: str) -> None:
    try:
        subprocess.run(
            command, check=True, capture_output=True, text=True, encoding="utf-8", errors="replace"
        )
    except FileNotFoundError as exc:
        raise TrimError("ffmpeg was not found. Install FFmpeg and add it to PATH.") from exc
    except subprocess.CalledProcessError as exc:
        detail = (exc.stderr or "").strip()
        raise TrimError(f"{action}.{(' ' + detail) if detail else ''}") from exc
    except OSError as exc:
        raise TrimError(f"Could not run ffmpeg: {exc}") from exc


def trim_proxy(
    proxy_path: str | Path,
    workbook_path: str | Path,
    output_path: str | Path,
    *,
    include_audio: bool = True,
    overwrite: bool = False,
) -> TrimResult:
    """Join spreadsheet-selected ranges from a proxy into one MP4 clip."""
    source_path = Path(proxy_path).expanduser()
    workbook_file = Path(workbook_path).expanduser()
    destination = Path(output_path).expanduser()
    if not source_path.is_file():
        raise TrimError(f"Proxy video does not exist or is not a file: '{source_path}'.")
    if not workbook_file.is_file():
        raise TrimError(f"Range workbook does not exist or is not a file: '{workbook_file}'.")
    if workbook_file.suffix.lower() != ".xlsx":
        raise TrimError("Range workbook must be an .xlsx file.")
    if destination.suffix.lower() != ".mp4":
        raise TrimError("Output path must use the .mp4 extension.")
    try:
        source_resolved = source_path.resolve(strict=True)
        workbook_resolved = workbook_file.resolve(strict=True)
        output_resolved = destination.resolve(strict=False)
    except OSError as exc:
        raise TrimError(f"Could not resolve one of the paths: {exc}") from exc
    if len({source_resolved, workbook_resolved, output_resolved}) != 3:
        raise TrimError("Proxy, workbook, and output must be three distinct files.")
    if destination.exists() and not overwrite:
        raise TrimError(f"Output already exists: '{destination}'. Use another path or set overwrite=True.")

    ffmpeg = shutil.which("ffmpeg")
    ffprobe = shutil.which("ffprobe")
    missing = [name for name, path in (("ffmpeg", ffmpeg), ("ffprobe", ffprobe)) if path is None]
    if missing:
        raise TrimError(
            f"Required executable(s) not found: {', '.join(missing)}. "
            "Install FFmpeg (including ffprobe) and add both tools to PATH."
        )
    assert ffmpeg is not None and ffprobe is not None

    source_info = _probe_video(source_path, ffprobe)
    segments = read_segments(workbook_file, source_info.duration_seconds, source_info.frame_rate)
    use_audio = include_audio and source_info.has_audio
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise TrimError(f"Could not create output directory '{destination.parent}': {exc}") from exc

    temporary_path: Path | None = None
    try:
        handle = tempfile.NamedTemporaryFile(
            prefix=f".{destination.stem}-", suffix=".tmp.mp4", dir=destination.parent, delete=False
        )
        temporary_path = Path(handle.name)
        handle.close()
        temporary_path.unlink()

        graph = _filter_graph(segments, use_audio)
        command = [
            ffmpeg,
            "-hide_banner", "-v", "error", "-nostdin", "-y",
            "-i", str(source_path),
            "-filter_complex", graph,
            "-map", "[vout]",
        ]
        if use_audio:
            command += ["-map", "[aout]", "-c:a", "aac", "-b:a", "128k"]
        else:
            command += ["-an"]
        command += [
            "-c:v", "libx264", "-preset", "fast", "-crf", "23",
            "-pix_fmt", "yuv420p", "-fps_mode", "cfr",
            "-map_metadata", "-1", "-metadata:s:v:0", "rotate=0",
            "-movflags", "+faststart", str(temporary_path),
        ]
        _run(command, "FFmpeg could not create the trimmed proxy")

        if not temporary_path.is_file() or temporary_path.stat().st_size == 0:
            raise TrimError("FFmpeg completed without creating a usable output file.")
        output_info = _probe_video(temporary_path, ffprobe)
        if (output_info.display_width, output_info.display_height) != (
            source_info.display_width, source_info.display_height
        ):
            raise TrimError("The trimmed clip's display dimensions do not match the source proxy.")
        if abs(output_info.frame_rate - source_info.frame_rate) > max(0.01, source_info.frame_rate * 0.001):
            raise TrimError("The trimmed clip's frame rate does not match the source proxy.")
        if use_audio and not output_info.has_audio:
            raise TrimError("The source proxy has audio, but the trimmed clip does not.")

        try:
            if overwrite:
                os.replace(temporary_path, destination)
            else:
                os.link(temporary_path, destination)
                temporary_path.unlink()
        except FileExistsError as exc:
            raise TrimError(f"Output already exists: '{destination}'. Use another path or set overwrite=True.") from exc
        except OSError as exc:
            raise TrimError(f"Could not publish trimmed clip to '{destination}': {exc}") from exc
        temporary_path = None
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass

    kept_duration = sum(segment.frame_count for segment in segments) / source_info.frame_rate
    return TrimResult(
        source=source_info,
        output=output_info,
        segments=segments,
        output_path=destination,
        kept_duration_seconds=kept_duration,
        audio_included=use_audio,
    )


def _format_seconds(value: float) -> str:
    return f"{value:.3f} s"


def _print_result(result: TrimResult) -> None:
    print(f"Proxy: {result.source.display_width}x{result.source.display_height}, "
          f"{result.source.frame_rate:.3f} fps, duration {_format_seconds(result.source.duration_seconds)}, "
          f"audio {'yes' if result.source.has_audio else 'no'}")
    for number, segment in enumerate(result.segments, start=1):
        label = f" ({segment.label})" if segment.label else ""
        print(
            f"Range {number}{label}, workbook row {segment.row_number}: "
            f"requested {_format_seconds(segment.start_seconds)}–{_format_seconds(segment.stop_seconds)}; "
            f"effective {_format_seconds(segment.effective_start_seconds)}–"
            f"{_format_seconds(segment.effective_stop_seconds)}"
        )
    print(f"Kept duration: {_format_seconds(result.kept_duration_seconds)}")
    print(f"Output: {result.output.display_width}x{result.output.display_height}, "
          f"{result.output.frame_rate:.3f} fps, duration {_format_seconds(result.output.duration_seconds)}, "
          f"audio {'yes' if result.output.has_audio else 'no'}")
    print(f"Output path: {result.output_path}")


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Keep selected time ranges from a proxy video using a Segments .xlsx workbook."
    )
    parser.add_argument("proxy", type=Path, help="Step 1 proxy MP4")
    parser.add_argument("ranges", type=Path, help=".xlsx workbook containing the Segments worksheet")
    parser.add_argument("output", type=Path, help="combined trimmed output path (.mp4)")
    parser.add_argument("--no-audio", action="store_true", help="remove audio from the output clip")
    parser.add_argument("--overwrite", action="store_true", help="replace an existing output file")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    try:
        result = trim_proxy(
            args.proxy,
            args.ranges,
            args.output,
            include_audio=not args.no_audio,
            overwrite=args.overwrite,
        )
    except TrimError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    _print_result(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
