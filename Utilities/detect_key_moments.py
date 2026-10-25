"""Find additional sport-specific key-moment candidates in a trimmed video.

This first-pass detector uses deterministic frame-to-frame motion and visual
change cues. It does not classify events or predict labels. Install Python
dependencies with ``python -m pip install -r Utilities/requirements.txt`` and
install FFmpeg/ffprobe separately on PATH.

CLI::

    python Utilities/detect_key_moments.py output/clips/example_selected.mp4 \
        --sport soccer --examples Test/key_moments_example.xlsx
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import Any


class DetectionError(RuntimeError):
    """An input, runtime, detection, or output error with user-facing context."""


class SportType(str, Enum):
    SOCCER = "soccer"
    BASKETBALL = "basketball"
    HOCKEY = "hockey"


@dataclass(frozen=True)
class SportProfile:
    event_window_seconds: float
    sample_rate_hz: float
    threshold: float
    suppression_seconds: float
    example_calibration_fraction: float
    motion_weight: float
    change_weight: float


# Profiles are intentionally simple and visible: later sport-specific models
# can replace cues without changing workbook, scan, or publishing contracts.
SPORT_PROFILES: dict[SportType, SportProfile] = {
    SportType.SOCCER: SportProfile(2.0, 5.0, 0.18, 2.0, 0.45, 0.70, 0.30),
    SportType.BASKETBALL: SportProfile(1.4, 6.0, 0.20, 1.5, 0.45, 0.65, 0.35),
    SportType.HOCKEY: SportProfile(1.6, 6.0, 0.21, 1.7, 0.45, 0.72, 0.28),
}


@dataclass(frozen=True)
class VideoInfo:
    duration_seconds: float
    fps: float | None
    width: int
    height: int


@dataclass(frozen=True)
class Annotation:
    time_seconds: float
    label: str
    row_number: int


@dataclass(frozen=True)
class Detection:
    time_seconds: float
    score: float


@dataclass(frozen=True)
class DetectionResult:
    video_path: Path
    sport: SportType
    example_count: int
    detections: tuple[Detection, ...]
    output_path: Path
    video: VideoInfo
    profile: SportProfile
    threshold: float


def _seconds(value: Any, row: int) -> float:
    if value is None or isinstance(value, bool):
        raise DetectionError(f"Row {row}: Time (seconds) must be a nonnegative time value.")
    if isinstance(value, dt.timedelta):
        result = value.total_seconds()
    elif isinstance(value, dt.time):
        result = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    elif isinstance(value, (int, float)):
        result = float(value)
    elif isinstance(value, str):
        text = value.strip()
        time_match = re.fullmatch(r"(\d+):([0-5]\d):([0-5]\d)(?:\.(\d{1,6}))?", text)
        if time_match:
            hours, minutes, seconds, fraction = time_match.groups()
            result = int(hours) * 3600 + int(minutes) * 60 + int(seconds)
            if fraction:
                result += float(f"0.{fraction}")
        else:
            try:
                result = float(text)
            except ValueError as exc:
                raise DetectionError(
                    f"Row {row}: Time (seconds) must be decimal seconds or HH:MM:SS[.mmm]."
                ) from exc
    else:
        raise DetectionError(f"Row {row}: unsupported Time (seconds) value.")
    if not math.isfinite(result) or result < 0:
        raise DetectionError(f"Row {row}: Time (seconds) must be finite and nonnegative.")
    return float(result)


def read_annotations(path: str | Path, duration_seconds: float) -> tuple[Annotation, ...]:
    """Read and strictly validate the KeyMoments workbook input contract."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DetectionError(
            "Reading .xlsx inputs requires openpyxl. Install with "
            "'python -m pip install -r Utilities/requirements.txt'."
        ) from exc
    workbook_path = Path(path).expanduser()
    if not workbook_path.is_file() or workbook_path.suffix.lower() != ".xlsx":
        raise DetectionError(f"Examples workbook must be an existing .xlsx file: '{workbook_path}'.")
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    except Exception as exc:
        raise DetectionError(f"Could not read examples workbook '{workbook_path}': {exc}") from exc
    try:
        if "KeyMoments" not in workbook.sheetnames:
            raise DetectionError("Examples workbook must contain a worksheet named 'KeyMoments'.")
        sheet = workbook["KeyMoments"]
        headers = [cell.value.strip() if isinstance(cell.value, str) else cell.value
                   for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())]
        try:
            time_col = headers.index("Time (seconds)")
            label_col = headers.index("Label")
        except ValueError as exc:
            raise DetectionError(
                "The 'KeyMoments' worksheet must have 'Time (seconds)' and 'Label' headers in row 1."
            ) from exc
        annotations: list[Annotation] = []
        seen: set[float] = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            time_cell = cells[time_col] if time_col < len(cells) else None
            label_cell = cells[label_col] if label_col < len(cells) else None
            time_value = time_cell.value if time_cell else None
            label_value = label_cell.value if label_cell else None
            if time_value is None and (label_value is None or (isinstance(label_value, str) and not label_value.strip())):
                continue
            if (time_cell and time_cell.data_type == "f") or (label_cell and label_cell.data_type == "f"):
                raise DetectionError(f"Row {row_number}: formulas are not allowed; use literal values.")
            seconds = _seconds(time_value, row_number)
            if seconds > duration_seconds:
                raise DetectionError(
                    f"Row {row_number}: time {seconds:g}s exceeds clip duration {duration_seconds:.3f}s."
                )
            if not isinstance(label_value, str) or not label_value.strip():
                raise DetectionError(f"Row {row_number}: Label must be nonempty text.")
            if seconds in seen:
                raise DetectionError(f"Row {row_number}: duplicate annotation time {seconds:g}s.")
            seen.add(seconds)
            annotations.append(Annotation(seconds, label_value.strip(), row_number))
        if not annotations:
            raise DetectionError("The KeyMoments worksheet contains no valid examples.")
        return tuple(annotations)
    finally:
        workbook.close()


def _probe_video(path: Path) -> VideoInfo:
    ffprobe = shutil.which("ffprobe")
    if not ffprobe:
        raise DetectionError("ffprobe was not found. Install FFmpeg (including ffprobe) and add it to PATH.")
    try:
        proc = subprocess.run(
            [ffprobe, "-v", "error", "-select_streams", "v:0", "-show_entries",
             "stream=width,height,avg_frame_rate,r_frame_rate,duration:format=duration", "-of", "json", str(path)],
            check=True, capture_output=True, text=True, encoding="utf-8", errors="replace",
        )
        data = json.loads(proc.stdout)
        stream = data["streams"][0]
        stream_duration = stream.get("duration")
        try:
            duration = float(stream_duration) if stream_duration not in (None, "", "N/A") else math.nan
        except (TypeError, ValueError):
            duration = math.nan
        if not math.isfinite(duration):
            duration = float(data.get("format", {}).get("duration"))
        rate = stream.get("avg_frame_rate") or stream.get("r_frame_rate")
        fps = None
        if rate and rate != "0/0":
            numerator, denominator = map(float, rate.split("/"))
            fps = numerator / denominator if denominator else None
        info = VideoInfo(duration, fps, int(stream["width"]), int(stream["height"]))
    except (OSError, subprocess.CalledProcessError, ValueError, KeyError, IndexError, TypeError, json.JSONDecodeError) as exc:
        details = getattr(exc, "stderr", "")
        raise DetectionError(f"Could not probe video '{path}'. {details or exc}") from exc
    if not math.isfinite(info.duration_seconds) or info.duration_seconds <= 0:
        raise DetectionError("Video has no known positive duration.")
    return info


def _opencv():
    try:
        import cv2
    except ImportError as exc:
        raise DetectionError(
            "Video scanning requires OpenCV. Install Python dependencies with "
            "'python -m pip install -r Utilities/requirements.txt'."
        ) from exc
    return cv2


def _scan(video_path: Path, video: VideoInfo, annotations: tuple[Annotation, ...],
          profile: SportProfile) -> tuple[list[tuple[float, float]], float]:
    """Return (timestamp, score) samples and an example-calibrated threshold."""
    cv2 = _opencv()
    try:
        capture = cv2.VideoCapture(str(video_path))
    except cv2.error as exc:
        raise DetectionError(f"Could not initialize video decoding for '{video_path}': {exc}") from exc
    if not capture.isOpened():
        capture.release()
        raise DetectionError(f"OpenCV could not decode video '{video_path}'. Check its codec and file integrity.")
    step = 1.0 / profile.sample_rate_hz
    samples: list[tuple[float, float, float]] = []
    previous = None
    last_sample = -step
    last_frame_timestamp = 0.0
    try:
        while True:
            ok, frame = capture.read()
            if not ok:
                break
            timestamp = capture.get(cv2.CAP_PROP_POS_MSEC) / 1000.0
            if not math.isfinite(timestamp) or timestamp < 0:
                raise DetectionError("Video decoder returned an invalid frame timestamp.")
            last_frame_timestamp = timestamp
            if timestamp - last_sample + 1e-6 < step:
                continue
            last_sample = timestamp
            height, width = frame.shape[:2]
            scale = min(1.0, 640 / max(width, height))
            if scale < 1:
                frame = cv2.resize(frame, (max(1, int(width * scale)), max(1, int(height * scale))))
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            if previous is None:
                motion = change = 0.0
            else:
                flow = cv2.calcOpticalFlowFarneback(previous, gray, None, 0.5, 2, 15, 2, 5, 1.1, 0)
                magnitude = cv2.magnitude(flow[..., 0], flow[..., 1])
                motion = min(1.0, float(magnitude.mean()) / 12.0)
                change = min(1.0, float(cv2.absdiff(previous, gray).mean()) / 48.0)
            samples.append((timestamp, motion, change))
            previous = gray
    except cv2.error as exc:
        raise DetectionError(f"Video decoding failed for '{video_path}': {exc}") from exc
    finally:
        capture.release()
    if not samples:
        raise DetectionError("No decodable video frames were found.")
    if last_frame_timestamp < video.duration_seconds - max(0.5, 2 * step):
        raise DetectionError(
            f"Video decoding stopped at {last_frame_timestamp:.3f}s before the probed duration "
            f"of {video.duration_seconds:.3f}s; check the file for decode errors."
        )

    scores = [(t, profile.motion_weight * motion + profile.change_weight * change)
              for t, motion, change in samples]
    positive_scores = [max((score for t, score in scores if abs(t - a.time_seconds) <= profile.event_window_seconds / 2), default=0.0)
                       for a in annotations]
    # Examples calibrate a reproducible threshold; the lower bound remains
    # sport-profile specific so still/poor examples cannot make noise dominate.
    example_floor = sorted(positive_scores)[max(0, math.ceil(len(positive_scores) * 0.25) - 1)]
    threshold = max(profile.threshold, example_floor * profile.example_calibration_fraction)
    return scores, threshold


def _select_candidates(scores: list[tuple[float, float]], annotations: tuple[Annotation, ...],
                       profile: SportProfile, threshold: float, duration: float) -> tuple[Detection, ...]:
    radius = profile.event_window_seconds / 2
    peaks = []
    for index, (time_value, score) in enumerate(scores):
        if score < threshold:
            continue
        left = scores[index - 1][1] if index else -math.inf
        right = scores[index + 1][1] if index + 1 < len(scores) else -math.inf
        if score >= left and score >= right:
            event_time = min(duration, max(0.0, time_value))
            if all(abs(event_time - known.time_seconds) > radius for known in annotations):
                peaks.append(Detection(event_time, min(1.0, max(0.0, score))))
    merged: list[Detection] = []
    for candidate in sorted(peaks, key=lambda item: item.time_seconds):
        if merged and candidate.time_seconds - merged[-1].time_seconds <= profile.suppression_seconds:
            if candidate.score > merged[-1].score:
                merged[-1] = candidate
        else:
            merged.append(candidate)
    return tuple(merged)


def _write_workbook(path: Path, detections: tuple[Detection, ...], overwrite: bool) -> None:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise DetectionError("Output workbook requires openpyxl; install Utilities/requirements.txt.") from exc
    if path.exists() and not overwrite:
        raise DetectionError(f"Output already exists: '{path}'. Choose another path or pass --overwrite.")
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=".tmp.xlsx", dir=path.parent, delete=False) as temp:
            temp_path = Path(temp.name)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Detected Moments"
        sheet.append(["Time (seconds)", "Confidence (0-1)"])
        for detection in detections:
            sheet.append([round(detection.time_seconds, 3), round(detection.score, 4)])
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 22
        sheet["A1"].comment = None
        sheet["B1"].comment = __import__("openpyxl").comments.Comment(
            "Heuristic score: weighted normalized optical-flow motion and frame difference, range 0-1. It is not a probability or event classification.",
            "SportVideoAnalysis",
        )
        for cell in sheet[1]:
            cell.font = __import__("openpyxl").styles.Font(bold=True, color="FFFFFF")
            cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="17365D")
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, 1).number_format = "0.000"
            sheet.cell(row, 2).number_format = "0.0000"
        workbook.save(temp_path)
        workbook.close()
        check = load_workbook(temp_path, read_only=True, data_only=False)
        try:
            if "Detected Moments" not in check.sheetnames or list(next(check["Detected Moments"].iter_rows(min_row=1, max_row=1, values_only=True))) != ["Time (seconds)", "Confidence (0-1)"]:
                raise DetectionError("Generated results workbook failed verification.")
        finally:
            check.close()
        if overwrite:
            os.replace(temp_path, path)
        else:
            os.link(temp_path, path)
            temp_path.unlink()
        temp_path = None
    except FileExistsError as exc:
        raise DetectionError(f"Output already exists: '{path}'. Choose another path or pass --overwrite.") from exc
    except DetectionError:
        raise
    except Exception as exc:
        raise DetectionError(f"Could not write results workbook '{path}': {exc}") from exc
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def detect_key_moments(video_path: str | Path, sport: SportType | str, examples_path: str | Path,
                       output_path: str | Path | None = None, *, overwrite: bool = False) -> DetectionResult:
    """Scan a trimmed clip and write additional unlabeled candidate timestamps."""
    try:
        selected_sport = sport if isinstance(sport, SportType) else SportType(sport)
    except ValueError as exc:
        raise DetectionError(f"Unsupported sport '{sport}'. Supported values: {', '.join(s.value for s in SportType)}.") from exc
    video_file = Path(video_path).expanduser()
    examples_file = Path(examples_path).expanduser()
    if not video_file.is_file():
        raise DetectionError(f"Input video does not exist or is not a file: '{video_file}'.")
    video = _probe_video(video_file)
    annotations = read_annotations(examples_file, video.duration_seconds)
    profile = SPORT_PROFILES[selected_sport]
    scores, threshold = _scan(video_file, video, annotations, profile)
    detections = _select_candidates(scores, annotations, profile, threshold, video.duration_seconds)
    destination = Path(output_path).expanduser() if output_path else video_file.with_name(f"{video_file.stem}_key_moments.xlsx")
    if destination.suffix.lower() != ".xlsx":
        raise DetectionError("Output path must use the .xlsx extension.")
    if destination.resolve(strict=False) in {video_file.resolve(), examples_file.resolve()}:
        raise DetectionError("Output must be separate from the input video and annotation workbook.")
    if destination.exists() and not overwrite:
        raise DetectionError(f"Output already exists: '{destination}'. Choose another path or pass --overwrite.")
    _write_workbook(destination, detections, overwrite)
    return DetectionResult(video_file, selected_sport, len(annotations), detections, destination,
                           video, profile, threshold)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Find additional unlabeled key-moment candidates in a trimmed video.")
    parser.add_argument("video", type=Path, help="Step 2 trimmed clip, normally under output/clips/")
    parser.add_argument("--sport", required=True, choices=[sport.value for sport in SportType])
    parser.add_argument("--examples", required=True, type=Path, help=".xlsx workbook with KeyMoments annotations")
    parser.add_argument("--output", type=Path, help="results .xlsx (default: <clip-stem>_key_moments.xlsx beside video)")
    parser.add_argument("--overwrite", action="store_true", help="replace an existing results workbook")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        result = detect_key_moments(args.video, args.sport, args.examples, args.output, overwrite=args.overwrite)
    except DetectionError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    print(f"Video: {result.video_path} ({result.video.width}x{result.video.height}, {result.video.duration_seconds:.3f}s)")
    print(f"Sport profile: {result.sport.value}; examples: {result.example_count}; detections: {len(result.detections)}")
    print(f"Settings: sample={result.profile.sample_rate_hz:g}Hz, event window={result.profile.event_window_seconds:g}s, "
          f"threshold={result.threshold:.4f}, suppression={result.profile.suppression_seconds:g}s; "
          "timestamps rounded to nearest millisecond; confidence is a heuristic score in [0,1].")
    print(f"Output: {result.output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
